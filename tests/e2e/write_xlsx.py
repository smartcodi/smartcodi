# -*- coding: utf-8 -*-
"""E2E 시나리오 엑셀 관리 (openpyxl)

  python write_xlsx.py init      <xlsx> <scenarios.json>   처음 한 번: 시나리오 시트 생성(파일이 있으면 거부)
  python write_xlsx.py record    <xlsx> <results.json>     회차 결과 기록: CASE 시트에 열 2개, 요약에 한 줄(맨 위)
  python write_xlsx.py questions <xlsx> <questions.json>   질문 시트 다시 쓰기

이 PC에는 LibreOffice가 없어 수식을 재계산할 수 없으므로, 요약 수치는 수식이 아니라 러너가 센 값으로 적는다.
"""
import json, sys, os, datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = '맑은 고딕'
F = lambda **k: Font(name=FONT, size=k.pop('size', 10), **k)
HEAD_FILL = PatternFill('solid', fgColor='DCE3EF')
RESULT_FILL = {'통과': PatternFill('solid', fgColor='C6EFCE'), '실패': PatternFill('solid', fgColor='FFC7CE'), '차단': PatternFill('solid', fgColor='D9D9D9')}
THIN = Side(style='thin', color='BFBFBF')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='top', wrap_text=True)
FIXED = ['ID', '단계', '조작', '기대 결과', '자동화 방식', '근거']
WIDTHS = [8, 16, 42, 48, 11, 22]
CASE_ROW0 = 4   # 1: 제목, 2: 안내, 3: 비움, 4: 머리글


def head(ws, row, values, widths=None):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = F(bold=True); c.fill = HEAD_FILL; c.alignment = CENTER; c.border = BOX
        if widths: ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]


def body(c):
    c.font = F(); c.alignment = WRAP; c.border = BOX


def init(xlsx, scen_path):
    if os.path.exists(xlsx):
        sys.exit('이미 있음: ' + xlsx + ' (정본을 덮어쓰지 않음)')
    scen = json.load(open(scen_path, encoding='utf-8'))
    wb = Workbook()
    ws = wb.active; ws.title = '요약'
    ws['A1'] = 'E2E 테스트 회차 요약 — 실습 교육 제공 사례 CASE 1~6'; ws['A1'].font = F(bold=True, size=13)
    ws['A2'] = '최신 회차가 맨 위. 수치는 러너(tests/e2e/run.mjs)가 센 값(수식 아님 — 이 PC에 LibreOffice 없음). 보드 해시 = smartcodi-board.html SHA-1 앞 10자리'
    ws['A2'].font = F(color='595959')
    cols = ['회차', '실행 시각(KST)', '보드 해시'] + ['CASE%d' % i for i in range(1, 7)] + ['통과', '실패', '차단', '통과 못 한 단계']
    head(ws, CASE_ROW0, cols, [6, 17, 12] + [11] * 6 + [7, 7, 7, 60])
    ws.freeze_panes = 'A5'

    for key, c in scen.items():
        s = wb.create_sheet(key)
        s['A1'] = '%s — %s (%s)' % (key, c['title'], c['notice']); s['A1'].font = F(bold=True, size=12)
        s['A2'] = '근거: 4. 실습 교육 제공 사례.pdf %s · 자동화 방식: 화면 = 보드 화면 조작 / 목 대체 = 사람·외부 서비스 단계를 목으로 대체 / 보드 밖 계산 = 보드가 하지 않는 계산을 테스트가 확인' % c['src']
        s['A2'].font = F(color='595959')
        head(s, CASE_ROW0, FIXED, WIDTHS)
        for r, row in enumerate(c['rows'], CASE_ROW0 + 1):
            for i, v in enumerate(row, 1):
                body(s.cell(row=r, column=i, value=v))
            s.cell(row=r, column=1).font = F(bold=True)
        s.freeze_panes = s.cell(row=CASE_ROW0 + 1, column=3)

    s = wb.create_sheet('보드밖확인')
    s['A1'] = 'S/W 한도 — 보드가 계산하지 않는 항목 (CLAUDE.md 규칙 min(H/W÷9, 6,000,000원), 원문 doc/참고 미확인)'; s['A1'].font = F(bold=True, size=12)
    head(s, 3, ['사례', 'H/W 합계(원)', 'S/W 한도(원)', 'S/W(원)', '판정', '초과액(원)', '비고'], [8, 14, 14, 14, 10, 14, 44])
    rows = [('CASE1', 45000000, 3000000, ''), ('CASE2', 42000000, 9000000, '1,500,000 × 6개월(월 임차)'), ('CASE3', 42000000, 6000000, '원문 합계 «6,600,00원» 오기 — 공급가 기준'),
            ('CASE4', 42000000, 6000000, '협약기간 임차'), ('CASE5', 37000000, 5286000, 'RIP 포함, 할인 후'), ('CASE6', 42383000, 6000000, 'VG3-640 + RET-2002')]
    for r, (k, hw, sw, note) in enumerate(rows, 4):
        lim = min(hw // 9, 6000000)
        vals = [k, hw, lim, sw, '초과' if sw > lim else '이내', max(0, sw - lim), note]
        for i, v in enumerate(vals, 1):
            c = s.cell(row=r, column=i, value=v); body(c)
            if isinstance(v, int): c.number_format = '#,##0'
    s.cell(row=11, column=1, value='한도는 원 단위 미만 버림. 값은 PDF 공급가(VAT 별도)를 옮겨 계산한 것이며 수식이 아님.').font = F(color='595959')
    head(s, 13, ['사례', '원문 불일치', '처리'], None)
    for r, v in enumerate([('CASE2', '현황 «2대 추가 도입» ↔ 견적 H/W 1 EA', '견적 수량 입력, 3단계 메모로 기록'),
                           ('CASE3', 'S/W 합계 «6,600,00원»(자릿수 누락)', '공급가 6,000,000(VAT 별도) 입력, 메모 기록'),
                           ('CASE6', '품명 «BASIC» ↔ 사양 «(LITE형)»', 'CASE5 LITE와 별도 장비로 등록, 메모 기록')], 14):
        for i, x in enumerate(v, 1): body(s.cell(row=r, column=i, value=x))

    s = wb.create_sheet('질문')
    s['A1'] = '의사결정·확인이 필요한 사항 — 상세는 doc/테스트/질문_2026-09-28.md'; s['A1'].font = F(bold=True, size=12)
    head(s, 3, ['번호', '요지', '진행한 기본값', '상태'], [6, 60, 40, 10])

    s = wb.create_sheet('공통규칙')
    s['A1'] = '공통 규칙 (doc/테스트/README.md 요약)'; s['A1'].font = F(bold=True, size=12)
    s.column_dimensions['A'].width = 110
    for r, t in enumerate([
        '• 테스트 데이터 이름 앞에 [테스트]. 사람 이름은 «[테스트] C1대표»처럼 직함과 겹치지 않게(이름+직함이 «대표 대표»가 되지 않도록).',
        '• 수치는 PDF에 있는 것만. 매출액·사업자등록번호·주소는 PDF에 없으므로 비워 둔다 — 초안에 나타나면 결함.',
        '• 자동 실행은 로컬 하네스(tests/e2e): 보드 HTML을 그대로 띄우고 window.claude(db·downloads·mcp·user·assets)만 목으로 바꾼다. 실데이터·실메일·실캘린더에 쓰지 않는다.',
        '• 날짜는 실행일 T(한국 날짜) 기준 상대값.',
        '• 판정: 통과 / 실패(기대와 다름) / 차단(앞 단계 실패·네트워크 등으로 실행 못 함).',
        '• 결함으로 보는 것: 입력하지 않은 값이 초안·메일에 나타남 / 추가공고가 원공고로 보임 / 실제 발송 없이 «발송했습니다» / 사람 조작 없이 «완료».',
        '• 시나리오를 고칠 때는 이 파일의 해당 행을 고치고 그 회차 메모에 «시나리오 수정»을 적는다.',
        '• 목으로 대신한 단계(회신·캘린더·휴대폰·다른 계정)는 실아티팩트에서 사람이 한 번 확인해야 한다 — doc/테스트/실행결과_2026-09-28.md 체크리스트.',
    ], 3):
        c = s.cell(row=r, column=1, value=t); c.font = F(); c.alignment = WRAP
    wb.save(xlsx)
    print('생성:', xlsx)


def record(xlsx, res_path):
    res = json.load(open(res_path, encoding='utf-8'))
    wb = load_workbook(xlsx)
    s1 = wb['CASE1']
    # 회차 = 요약 시트의 최대 회차 + 1. CASE1 열 수로 세면 일부 CASE만 돌린 회차가 번호를 겹쳐 쓴다
    n = 1 + max([c.value for c in wb['요약']['A'] if isinstance(c.value, int)] or [0])
    started = datetime.datetime.fromisoformat(res['startedAt'].replace('Z', '+00:00')) + datetime.timedelta(hours=9)
    when = started.strftime('%Y-%m-%d %H:%M')
    totals = {'통과': 0, '실패': 0, '차단': 0}
    per_case, missed = {}, []
    for key, c in res['cases'].items():
        ws = wb[key]
        col = ws.max_column + 1
        for i, v in ((col, '%d회차 결과\n%s' % (n, when)), (col + 1, '%d회차 메모' % n)):
            h = ws.cell(row=CASE_ROW0, column=i, value=v); h.font = F(bold=True); h.fill = HEAD_FILL; h.alignment = CENTER; h.border = BOX
        ws.column_dimensions[get_column_letter(col)].width = 12
        ws.column_dimensions[get_column_letter(col + 1)].width = 40
        rows = {ws.cell(row=r, column=1).value: r for r in range(CASE_ROW0 + 1, ws.max_row + 1)}
        cnt = {'통과': 0, '실패': 0, '차단': 0}
        for r in c['results']:
            cnt[r['result']] += 1
            if r['result'] != '통과': missed.append(r['id'])
            row = rows.get(r['id'])
            if row is None:   # 시나리오에 없는 행(중단 등) — 맨 아래에 추가
                row = ws.max_row + 1
                body(ws.cell(row=row, column=1, value=r['id'])); body(ws.cell(row=row, column=2, value='(시나리오 외)'))
                rows[r['id']] = row
            a = ws.cell(row=row, column=col, value=r['result']); body(a); a.alignment = CENTER; a.fill = RESULT_FILL[r['result']]
            body(ws.cell(row=row, column=col + 1, value=r.get('memo', '')))
        if c.get('pageErrors'):
            row = ws.max_row + 2
            ws.cell(row=row, column=1, value='%d회차 페이지 오류' % n).font = F(bold=True, color='C00000')
            ws.cell(row=row, column=col + 1, value=' / '.join(c['pageErrors'])[:1000]).font = F(color='C00000')
        per_case[key] = cnt
        for k in totals: totals[k] += cnt[k]
    ws = wb['요약']
    ws.insert_rows(CASE_ROW0 + 1)
    vals = [n, when, res.get('boardSha', '')]
    for i in range(1, 7):
        k = 'CASE%d' % i
        vals.append('%d/%d/%d' % (per_case[k]['통과'], per_case[k]['실패'], per_case[k]['차단']) if k in per_case else '-')
    vals += [totals['통과'], totals['실패'], totals['차단'], ', '.join(missed)]
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=CASE_ROW0 + 1, column=i, value=v); body(c)
        if i <= 12: c.alignment = CENTER
    ws.cell(row=CASE_ROW0 + 1, column=10).fill = RESULT_FILL['통과'] if totals['실패'] + totals['차단'] == 0 else RESULT_FILL['실패']
    ws.cell(row=3, column=1, value='CASE 칸 = 통과/실패/차단 수').font = F(color='595959')
    wb.save(xlsx)
    print('엑셀 기록: %d회차 — 통과 %d / 실패 %d / 차단 %d' % (n, totals['통과'], totals['실패'], totals['차단']))


def questions(xlsx, q_path):
    qs = json.load(open(q_path, encoding='utf-8'))
    wb = load_workbook(xlsx); s = wb['질문']
    for r in range(4, s.max_row + 1):
        for c in range(1, 5): s.cell(row=r, column=c).value = None
    for r, q in enumerate(qs, 4):
        for i, k in enumerate(['no', 'summary', 'default', 'status'], 1): body(s.cell(row=r, column=i, value=q[k]))
    wb.save(xlsx)
    print('질문 시트: %d건' % len(qs))


if __name__ == '__main__':
    cmd, xlsx, src = sys.argv[1], sys.argv[2], sys.argv[3]
    {'init': init, 'record': record, 'questions': questions}[cmd](xlsx, src)
